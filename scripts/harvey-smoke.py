"""Real Docker/core/tools/parsers with an explicitly scripted model double."""
import argparse
import json
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'src'))
from localauthor.harvey.runner import run_case


OFFICE="""python - <<'PY'
import json,socket
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
import subprocess
root=Path('/workspace/output')
doc=Document();doc.add_paragraph('ShopAir software fixture');doc.save(root/'brief.docx')
book=Workbook();book.active['A1']='ShopAir software fixture';book.save(root/'brief.xlsx')
slides=Presentation();slide=slides.slides.add_slide(slides.slide_layouts[1]);slide.shapes.title.text='ShopAir software fixture';slides.save(root/'brief.pptx')
subprocess.run(['libreoffice','-env:UserInstallation=file:///tmp/lo','--headless','--convert-to','pdf','--outdir',str(root),str(root/'brief.docx')],check=True,capture_output=True)
checks={}
for key,path in [('documents_readonly','/workspace/documents/forbidden'),('source_readonly','/opt/localauthor/forbidden'),('history_readonly','/workspace/history/forbidden')]:
 try:Path(path).write_text('forbidden');checks[key]=False
 except OSError:checks[key]=True
try:socket.create_connection(('1.1.1.1',443),timeout=1);checks['network_blocked']=False
except OSError:checks['network_blocked']=True
checks['rubric_hidden']=not Path('/workspace/task.json').exists()
assert all(checks.values()),checks
(root/'isolation.json').write_text(json.dumps(checks))
print(json.dumps(checks))
PY"""


class ScriptedEngine:
    spec=SimpleNamespace(output_tokens=128,context_tokens=4096)
    def __init__(self):
        calls=[('bash',{'command':OFFICE}),('glob',{'pattern':'**/*','path':'/workspace/documents'}),
               ('grep',{'pattern':'ShopAir','path':'/workspace/documents','output_mode':'content'}),
               ('grep',{'pattern':'ShopAir','path':'/workspace/documents/context.md','output_mode':'content'}),
               ('grep',{'pattern':'ShopAir','path':'/workspace/documents/context.md','glob':'*.pdf','output_mode':'content'}),
               ('read',{'file_path':'/workspace/documents/context.md'})]
        calls += [('read',{'file_path':'/workspace/output/brief.'+ext}) for ext in ['docx','xlsx','pptx','pdf']]
        calls += [('write',{'file_path':'response.md','content':'Unfinished'}),
                  ('edit',{'file_path':'response.md','old_string':'Unfinished','new_string':'Software flow validated','replace_all':False}),
                  ('finish',{'summary':'explicit software fixture finished','deliverables':['response.md']})]
        self.actions=iter(calls)
    def count(self,messages):return 128
    def generate(self,messages):
        if 'agent_output' in messages[-1]['content'] and 'match_criteria' in messages[-1]['content']:
            return json.dumps({'reasoning':'Explicit judge double; software only.','verdict':'pass'}),False
        try:name,arguments=next(self.actions)
        except StopIteration:
            if getattr(self,'natural_final',False):return json.dumps({'final':'scripted stop without finish'}),False
            raise
        return json.dumps({'tool':name,'arguments':arguments}),False


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--image',required=True);parser.add_argument('--report',type=Path,required=True)
    args=parser.parse_args()
    with tempfile.TemporaryDirectory(prefix='localauthor-harvey-') as directory:
        run=Path(directory);(run/'documents').mkdir();(run/'documents/context.md').write_text('ShopAir synthetic context',encoding='utf-8')
        task={'title':'Harvey software contract fixture','instructions':'Exercise sandbox tools. This is a scripted software fixture, not model qualification.',
              'deliverables':{'response.md':'response.md'},'criteria':[{'id':'C1','title':'Completed artifact','match_criteria':'response.md contains Software flow validated','deliverables':['response.md']}]}
        result=run_case(task,run,args.image,ScriptedEngine(),max_turns=16)
        transcript=[json.loads(line) for line in (run/'transcript.jsonl').read_text(encoding='utf-8').splitlines()]
        tools=[row for row in transcript if row['role']=='tool'];names={row['tool_name'] for row in tools}
        assert names=={'bash','read','write','edit','glob','grep','finish'},names
        for row in tools:
            assert not row['result'].startswith('Error:'),row
            if row['tool_name']=='read':assert 'ShopAir' in row['result'],row
            if row['tool_name']=='grep':
                parameters=json.loads(row['arguments'])
                if parameters.get('glob')=='*.pdf':assert row['result'].startswith('No matches'),row
                else:assert 'ShopAir' in row['result'] and 'No matches' not in row['result'],row
        assert (run/'output/response.md').read_text()=='Software flow validated'
        isolation=json.loads((run/'output/isolation.json').read_text());assert all(isolation.values())
        assert result['scores']['score']==1 and result['metrics']['finished_cleanly'] and not result['real_model']
        # A deliberately permissive judge must not hide unsupported features,
        # including an agent that skips finish and returns a natural final.
        negative=run/'natural-stop-guard';(negative/'documents').mkdir(parents=True)
        (negative/'documents/context.md').write_text('ShopAir oferece catálogo e organização de pedidos.',encoding='utf-8')
        engine=ScriptedEngine();engine.natural_final=True
        engine.actions=iter([('read',{'file_path':'/workspace/documents/context.md'}),
            ('write',{'file_path':'response.md','content':'ShopAir oferece estoque e processamento automatizado.'})])
        negative_task={**task,'title':'Marketing grounding negative fixture'}
        blocked=run_case(negative_task,negative,args.image,engine,max_turns=4)
        assert blocked['scores']['score']==1 and blocked['metrics']['finished_cleanly']
        assert set(blocked['grounding']['unsupported_families'])=={'estoque','automação'},blocked
        checks=[json.loads(line) for line in (negative/'grounding-checks.jsonl').read_text(encoding='utf-8').splitlines()]
        assert checks[-1]['decision']=='reject' and checks[-1]['candidate'],checks
        report={'success':True,'real_model':False,'judge':'explicit double','tools':sorted(names),'document_parsers':['docx','xlsx','pptx','pdf'],
                'isolation':isolation,'grounding_detects_finish_bypass':True,'upstream_loop':True,'upstream_scorer':True,'upstream_report':(run/'upstream-report.html.untrusted').exists()}
        args.report.parent.mkdir(parents=True,exist_ok=True);args.report.write_text(json.dumps(report,indent=2),encoding='utf-8');print(json.dumps(report))


if __name__=='__main__':main()
